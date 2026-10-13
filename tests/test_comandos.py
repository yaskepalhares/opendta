"""Fase 8: lowess, lrtest, hausman, xtpoisson, xtlogit, qreg e putexcel."""

import numpy as np
import pytest

DATA = """clear
set seed 804
set obs 200
gen id = ceil(_n/10)
gen t = mod(_n-1, 10) + 1
xtset id t
gen x = round(runiform()*4, .01)
gen u = (id - 10)/10
gen yc = floor(exp(.2 + .3*x + u) * runiform()*2)
gen yb = runiform() < invlogit(-1 + .6*x + u)
gen y = 1 + 2*x + u + round(runiform(), .01)
"""


def _c(run, n):
    return run.session.data.get(n).data.astype(float)


def test_lowess_full_window_without_weights_is_ols(run):
    run(DATA)
    run("lowess y x, bwidth(1) noweight generate(lw) nograph")
    x, y, lw = _c(run, "x"), _c(run, "y"), _c(run, "lw")
    b = np.polyfit(x, y, 1)
    assert np.allclose(lw, np.polyval(b, x), atol=1e-4)


def test_lrtest_and_hausman(run):
    run(DATA)
    run("quietly logit yb x\nestimates store a\nquietly logit yb x u\nestimates store b")
    run("lrtest b a")
    ll = {}
    run("quietly logit yb x")
    ll["a"] = run.session.e["ll"]
    run("quietly logit yb x u")
    ll["b"] = run.session.e["ll"]
    run("lrtest b a")
    assert run.session.r["chi2"] == pytest.approx(2 * (ll["b"] - ll["a"]), rel=1e-12)
    assert run.session.r["df"] == 1
    run("quietly xtreg y x, fe\nestimates store fe\nquietly xtreg y x, re\nestimates store re")
    run("hausman fe re")
    run("quietly xtreg y x, fe")
    b1, v1 = run.session.e["b"].data[0, 0], run.session.e["V"].data[0, 0]
    run("quietly xtreg y x, re")
    b2, v2 = run.session.e["b"].data[0, 0], run.session.e["V"].data[0, 0]
    run("hausman fe re")
    assert run.session.r["chi2"] == pytest.approx((b1 - b2) ** 2 / (v1 - v2), rel=1e-8)


def test_xtpoisson_fe_equals_poisson_with_dummies(run):
    run(DATA)
    run("quietly xtpoisson yc x, fe")
    b_fe = run.session.e["b"].data[0, 0]
    run("quietly poisson yc x i.id")
    b_d = run.session.e["b"].data[0, 0]
    assert b_fe == pytest.approx(b_d, rel=1e-6)


def test_xtlogit_fe_two_periods_equals_logit_on_differences(run):
    run("""clear
set seed 9
set obs 400
gen id = ceil(_n/2)
gen t = mod(_n-1, 2) + 1
xtset id t
gen x = runiform()*3
gen yb = runiform() < invlogit(-1 + .8*x + (id/200))""")
    run("quietly xtlogit yb x, fe")
    b_fe = run.session.e["b"].data[0, 0]
    ds = run.session.data
    ids, x, yb = _c(run, "id"), _c(run, "x"), _c(run, "yb")
    dx, dy = [], []
    for g in np.unique(ids):
        sel = ids == g
        if yb[sel].sum() == 1:
            xx, yy = x[sel], yb[sel]
            dx.append(xx[1] - xx[0])
            dy.append(yy[1])
    from scipy.optimize import minimize
    dx, dy = np.array(dx), np.array(dy)
    nll = lambda b: -np.sum(dy * b[0] * dx - np.log1p(np.exp(b[0] * dx)))   # noqa: E731
    b = minimize(nll, [0.0], tol=1e-12).x[0]
    assert b_fe == pytest.approx(b, rel=1e-5)


def test_qreg_median_is_lp_optimum(run):
    run(DATA)
    run("quietly qreg y x")
    b = run.session.e["b"].data.ravel()
    x, y = _c(run, "x"), _c(run, "y")

    def obj(bb):
        r = y - (bb[0] * x + bb[1])
        return np.sum(np.abs(r)) / 2
    base = obj(b)
    for d in ([1e-3, 0], [-1e-3, 0], [0, 1e-3], [0, -1e-3]):
        assert obj(b + np.array(d)) >= base - 1e-9
    assert run.session.e["sum_adev"] == pytest.approx(base, rel=1e-9)


def test_putexcel_writes_cells(run, tmp_path):
    import openpyxl
    run(DATA)
    f = tmp_path / "out.xlsx"
    run("quietly regress y x\nmatrix b = e(b)")
    out = run(f'putexcel set "{f}", replace\nputexcel A1 = "Titulo" B2 = (e(N)) A3 = matrix(b), names')
    assert "saved" in out
    ws = openpyxl.load_workbook(f).active
    assert ws["A1"].value == "Titulo"
    assert ws["B2"].value == 200
    assert ws["B3"].value == "x"
